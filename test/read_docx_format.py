from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os

def read_docx_format(docx_path):
    if not os.path.exists(docx_path):
        print(f"文件不存在: {docx_path}")
        return
    
    doc = Document(docx_path)
    
    print("=" * 60)
    print("Word 文档格式分析")
    print("=" * 60)
    
    # 页面设置
    section = doc.sections[0]
    print("\n【页面设置】")
    print(f"页面宽度: {section.page_width.inches} 英寸 ({section.page_width.cm} 厘米)")
    print(f"页面高度: {section.page_height.inches} 英寸 ({section.page_height.cm} 厘米)")
    print(f"左边距: {section.left_margin.inches} 英寸 ({section.left_margin.cm} 厘米)")
    print(f"右边距: {section.right_margin.inches} 英寸 ({section.right_margin.cm} 厘米)")
    print(f"上边距: {section.top_margin.inches} 英寸 ({section.top_margin.cm} 厘米)")
    print(f"下边距: {section.bottom_margin.inches} 英寸 ({section.bottom_margin.cm} 厘米)")
    print(f"页眉边距: {section.header_distance.inches} 英寸 ({section.header_distance.cm} 厘米)")
    print(f"页脚边距: {section.footer_distance.inches} 英寸 ({section.footer_distance.cm} 厘米)")
    
    # 页眉
    print("\n【页眉】")
    header = section.header
    if header.paragraphs:
        print(f"页眉段落数: {len(header.paragraphs)}")
        for i, para in enumerate(header.paragraphs):
            print(f"\n页眉段落 {i+1}:")
            print(f"  文本: {para.text}")
            print(f"  对齐方式: {para.alignment}")
            if para.runs:
                run = para.runs[0]
                print(f"  字体名称: {run.font.name}")
                if run.font.size:
                    print(f"  字体大小: {run.font.size.pt} 磅")
                print(f"  加粗: {run.font.bold}")
                print(f"  斜体: {run.font.italic}")
    
    # 页脚
    print("\n【页脚】")
    footer = section.footer
    if footer.paragraphs:
        print(f"页脚段落数: {len(footer.paragraphs)}")
        for i, para in enumerate(footer.paragraphs):
            print(f"\n页脚段落 {i+1}:")
            print(f"  文本: {para.text}")
            print(f"  对齐方式: {para.alignment}")
            if para.runs:
                run = para.runs[0]
                print(f"  字体名称: {run.font.name}")
                if run.font.size:
                    print(f"  字体大小: {run.font.size.pt} 磅")
                print(f"  加粗: {run.font.bold}")
                print(f"  斜体: {run.font.italic}")
    
    # 正文样式
    print("\n【正文样式】")
    if doc.paragraphs:
        # 分析前几个段落
        sample_count = min(5, len(doc.paragraphs))
        for i in range(sample_count):
            para = doc.paragraphs[i]
            if para.text.strip():
                print(f"\n正文段落 {i+1}:")
                print(f"  样式: {para.style.name}")
                print(f"  对齐方式: {para.alignment}")
                if para.runs:
                    run = para.runs[0]
                    print(f"  字体名称: {run.font.name}")
                    if run.font.size:
                        print(f"  字体大小: {run.font.size.pt} 磅")
                    print(f"  加粗: {run.font.bold}")
                    print(f"  斜体: {run.font.italic}")
    
    # 文档属性
    print("\n【文档属性】")
    print(f"标题: {doc.core_properties.title}")
    print(f"作者: {doc.core_properties.author}")
    print(f"主题: {doc.core_properties.subject}")
    print(f"关键词: {doc.core_properties.keywords}")
    
    print("\n" + "=" * 60)

if __name__ == "__main__":
    docx_path = r"d:\Junyu_DataStorage\01_Workspace\Tool_20260604_SoftwareCopyright-CodeDocGen\SC-CodeDocGen\docs\源程序.docx"
    read_docx_format(docx_path)
